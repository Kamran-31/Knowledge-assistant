"""
Extracts raw text from the supported file types:
PDF, DOCX, XLSX/XLS, CSV, TXT, and images (JPG/PNG via OCR).

Every parser returns a plain string. Callers should not need to know
anything about the underlying file format.
"""

from __future__ import annotations

import io

import fitz  # PyMuPDF
import pandas as pd
from PIL import Image
import pytesseract
from docx import Document


class UnsupportedFileTypeError(Exception):
    pass


def _extract_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text("text")
            if page_text.strip():
                text_parts.append(f"[Page {page_num}]\n{page_text}")
    return "\n\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of any tables in the document
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    return "\n".join(parts)


def _extract_excel(file_bytes: bytes, filename: str) -> str:
    engine = "openpyxl" if filename.lower().endswith(("xlsx", "xlsm")) else None
    sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, engine=engine)
    parts = []
    for sheet_name, df in sheets.items():
        parts.append(f"[Sheet: {sheet_name}]")
        parts.append(df.to_csv(index=False))
    return "\n\n".join(parts)


def _extract_csv(file_bytes: bytes) -> str:
    df = pd.read_csv(io.BytesIO(file_bytes))
    return df.to_csv(index=False)


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_image(file_bytes: bytes) -> str:
    image = Image.open(io.BytesIO(file_bytes))
    if image.mode != "RGB":
        image = image.convert("RGB")
    text = pytesseract.image_to_string(image)
    return text


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Dispatches to the right parser based on file extension.
    Returns extracted plain text (may be empty if nothing could be read).
    """
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext in ("xlsx", "xls"):
        return _extract_excel(file_bytes, filename)
    if ext == "csv":
        return _extract_csv(file_bytes)
    if ext == "txt":
        return _extract_txt(file_bytes)
    if ext in ("jpg", "jpeg", "png"):
        return _extract_image(file_bytes)

    raise UnsupportedFileTypeError(f"Unsupported file type: .{ext}")
